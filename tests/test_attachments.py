"""Локальные проверки вложений; не добавлять в коммит."""

import base64
import json
from contextlib import nullcontext
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app import agent, main
from app.attachments import ALLOWED_TYPES, MAX_FILE_SIZE, MAX_TEXT_LENGTH, extract_text

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j3ioAAAAASUVORK5CYII="
)


def xlsx_bytes(rows=None):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows or [("артикул", "количество"), ("001-ABC", 12), (None, None), ("XYZ-02", 3)]:
        sheet.append(row)
    workbook.create_sheet("Второй лист").append(["SECOND-SHEET", 7])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def pdf_bytes(page_count=1):
    writer = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    for index in range(page_count):
        page = writer.add_blank_page(width=300, height=300)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        })
        contents = DecodedStreamObject()
        contents.set_data(f"BT /F1 12 Tf 10 200 Td (PAGE-{index + 1}: SKU-123 quantity 5) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(contents)
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()


@pytest.fixture
def captured_agent(monkeypatch, no_openai):
    calls = []

    def fake_agent(message, history, product_ids, session_id, **kwargs):
        calls.append({"message": message, "history": history, "product_ids": product_ids, **kwargs})
        return {"reply": "Проверено", "products": []}

    monkeypatch.setattr(main, "run_agent", fake_agent)
    return calls


def test_xlsx_articles_quantities_and_all_sheets():
    text = extract_text(xlsx_bytes(), ".xlsx")
    assert text == "артикул | количество\n001-ABC | 12\nXYZ-02 | 3\nSECOND-SHEET | 7"


def test_pdf_extracts_real_text_and_only_first_ten_pages():
    assert "SKU-123 quantity 5" in extract_text(pdf_bytes(), ".pdf")
    text = extract_text(pdf_bytes(11), ".pdf")
    assert "PAGE-10:" in text
    assert "PAGE-11:" not in text
    assert "первые 10 страниц" in text


def test_docx_paragraphs_and_tables_in_document_order():
    document = Document()
    document.add_paragraph("Спецификация")
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["001-ABC", "Кабель", "12"]):
        cell.text = value
    document.add_paragraph("Конец")
    stream = BytesIO()
    document.save(stream)
    assert extract_text(stream.getvalue(), ".docx") == "Спецификация\n001-ABC | Кабель | 12\nКонец"


def test_document_text_is_limited_with_notice():
    text = extract_text(xlsx_bytes([("X" * 9000, 1)]), ".xlsx")
    assert len(text) == MAX_TEXT_LENGTH
    assert text.endswith("[Текст обрезан: лимит 8000 символов.]")


@pytest.mark.parametrize(("name", "mime"), [
    ("bad.txt", "text/plain"),
    ("bad.exe", "application/pdf"),
    ("bad.pdf", "image/jpeg"),
    ("bad.xlsx", "application/octet-stream"),
    ("old.doc", "application/msword"),
])
async def test_unsupported_type_returns_415_reply(client, captured_agent, name, mime):
    response = await client.post("/chat", data={"message": ""}, files={"files": (name, b"bad", mime)})
    assert response.status_code == 415
    assert "Неподдерживаемый тип файла" in response.json()["reply"]
    assert response.json()["products"] == []
    assert captured_agent == []


async def test_multipart_file_without_message_and_history(client, captured_agent):
    history = [{"role": "user", "content": "Нужен кабель"}]
    response = await client.post("/chat", data={
        "history": json.dumps(history), "product_ids": "[123]",
    }, files={"files": ("спецификация.xlsx", xlsx_bytes(), ALLOWED_TYPES[".xlsx"])})
    assert response.status_code == 200
    call = captured_agent[0]
    assert call["message"] == ""
    assert call["history"] == history
    assert call["product_ids"] == [123]
    assert "Содержимое вложения «спецификация.xlsx»:\nартикул | количество" in call["attachment_parts"][0]["text"]


async def test_attachment_over_1000_chars_is_allowed(client, captured_agent):
    response = await client.post("/chat", data={"message": "X" * 1000}, files={
        "files": ("large.xlsx", xlsx_bytes([("Y" * 3000, 1)]), ALLOWED_TYPES[".xlsx"]),
    })
    assert response.status_code == 200
    assert len(captured_agent[0]["message"]) == 1000
    assert len(captured_agent[0]["attachment_parts"][0]["text"]) > 3000


async def test_typed_message_over_1000_is_rejected(client, captured_agent):
    response = await client.post("/chat", data={"message": "X" * 1001}, files={
        "files": ("photo.png", PNG, "image/png"),
    })
    assert "1000" in response.json()["reply"]
    assert captured_agent == []


@pytest.mark.parametrize("count", [3, 4])
async def test_file_count_limit(client, captured_agent, count):
    response = await client.post("/chat", files=[
        ("files", (f"{index}.png", PNG, "image/png")) for index in range(count)
    ])
    if count == 3:
        assert response.status_code == 200
        assert len(captured_agent[0]["attachment_parts"]) == 6
    else:
        assert response.status_code == 413
        assert "не более 3" in response.json()["reply"]
        assert captured_agent == []


@pytest.mark.parametrize("extra", [0, 1])
async def test_file_size_limit(client, captured_agent, extra):
    data = PNG + b"\0" * (MAX_FILE_SIZE + extra - len(PNG))
    response = await client.post("/chat", files={"files": ("photo.png", data, "image/png")})
    if extra:
        assert response.status_code == 413
        assert "10 МБ" in response.json()["reply"]
        assert captured_agent == []
    else:
        assert response.status_code == 200
        assert captured_agent


async def test_oversized_upload_stops_reading_stream(client, captured_agent):
    consumed_tail = False

    async def body():
        nonlocal consumed_tail
        yield (
            b'--test-boundary\r\nContent-Disposition: form-data; name="files"; filename="photo.png"\r\n'
            b'Content-Type: image/png\r\n\r\n'
        )
        yield PNG + b"\0" * MAX_FILE_SIZE
        consumed_tail = True
        yield b"\r\n--test-boundary--\r\n"

    response = await client.post("/chat", content=body(), headers={
        "Content-Type": "multipart/form-data; boundary=test-boundary",
    })
    assert response.status_code == 413
    assert not consumed_tail
    assert captured_agent == []


@pytest.mark.parametrize("history", ["not JSON", "{}", '["bad entry"]'])
async def test_invalid_history_returns_reply(client, captured_agent, history):
    response = await client.post("/chat", data={"history": history}, files={"files": ("p.png", PNG, "image/png")})
    assert response.status_code == 422
    assert response.json()["reply"]
    assert captured_agent == []


@pytest.mark.parametrize(("name", "data", "mime"), [
    ("broken.pdf", b"not PDF", "application/pdf"),
    ("empty.pdf", b"", "application/pdf"),
    ("broken.xlsx", b"not a spreadsheet", ALLOWED_TYPES[".xlsx"]),
    ("broken.docx", b"not a document", ALLOWED_TYPES[".docx"]),
    ("broken.png", b"not an image", "image/png"),
])
async def test_unreadable_files_return_reply(client, captured_agent, name, data, mime):
    response = await client.post("/chat", files={"files": (name, data, mime)})
    assert response.status_code == 422
    assert name in response.json()["reply"]
    assert captured_agent == []


async def test_json_backward_compatibility_and_multipart_without_files(client, captured_agent):
    for kwargs in [
        {"json": {"message": "Кабель", "history": []}},
        {"files": {"message": (None, "Кабель"), "history": (None, "[]")}},
    ]:
        response = await client.post("/chat", **kwargs)
        assert response.status_code == 200
        assert response.json()["reply"] == "Проверено"
        assert captured_agent[-1]["message"] == "Кабель"
        assert "attachment_parts" not in captured_agent[-1]


async def test_empty_request_asks_for_input(client, captured_agent):
    response = await client.post("/chat", json={"message": "  "})
    assert response.status_code == 200
    assert "сформулируйте вопрос" in response.json()["reply"]
    assert captured_agent == []


async def test_photo_and_document_reach_model_as_user_content(client, monkeypatch):
    requests = []

    def create(**kwargs):
        requests.append(kwargs)
        message = SimpleNamespace(tool_calls=None, content="Вижу товар и спецификацию.")
        return SimpleNamespace(id="test", choices=[SimpleNamespace(message=message)])

    fake = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    monkeypatch.setattr(agent, "OpenAI", lambda **kwargs: nullcontext(fake))
    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    response = await client.post("/chat", data={"message": "Найди товар"}, files=[
        ("files", ("photo.PNG", PNG, "image/png")),
        ("files", ("invoice.pdf", pdf_bytes(), "application/pdf")),
    ])
    assert response.status_code == 200
    content = requests[0]["messages"][-1]["content"]
    assert content[0] == {"type": "text", "text": "Найди товар"}
    assert content[2] == {"type": "image_url", "image_url": {
        "url": "data:image/png;base64," + base64.b64encode(PNG).decode("ascii"),
    }}
    assert "Содержимое вложения «invoice.pdf»" in content[3]["text"]
    assert "SKU-123" in content[3]["text"]
