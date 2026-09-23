"""Проверка вложений и подготовка текстовых/визуальных частей сообщения."""

import base64
from collections.abc import Iterator
from contextlib import asynccontextmanager, closing
from io import BytesIO
from pathlib import PurePosixPath

import xlrd
from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from pypdf import PdfReader
from starlette.datastructures import UploadFile
from starlette.formparsers import MultiPartException, MultiPartParser
from starlette.requests import Request

MAX_FILES = 3
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_TEXT_LENGTH = 8000
TRUNCATION_NOTICE = "\n[Текст обрезан: лимит 8000 символов.]"
ALLOWED_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".pdf": "application/pdf",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class AttachmentError(ValueError):
    def __init__(self, reply: str, status_code: int = 422) -> None:
        super().__init__(reply)
        self.reply = reply
        self.status_code = status_code


class _FileTooLarge(MultiPartException):
    pass


class _AttachmentParser(MultiPartParser):
    def on_part_begin(self) -> None:
        super().on_part_begin()
        self._file_size = 0

    def on_part_data(self, data: bytes, start: int, end: int) -> None:
        if self._current_part.file is not None:
            self._file_size += end - start
            if self._file_size > MAX_FILE_SIZE:
                name = attachment_name(self._current_part.file.filename)
                # MultiPartParser закроет временные файлы при MultiPartException.
                raise _FileTooLarge(f"Файл «{name}» превышает лимит 10 МБ.")
        super().on_part_data(data, start, end)


@asynccontextmanager
async def attachment_form(request: Request):
    """Ограничивает файлы во время загрузки и всегда закрывает временные файлы."""
    parser = _AttachmentParser(request.headers, request.stream(), max_files=MAX_FILES, max_fields=3)
    try:
        form = await parser.parse()
    except _FileTooLarge as error:
        raise AttachmentError(error.message, 413) from error
    except MultiPartException as error:
        if "Too many files" in error.message:
            raise AttachmentError("Можно прикрепить не более 3 файлов.", 413) from error
        raise AttachmentError(
            "Не удалось прочитать форму. Проверьте поля message, history и files.", 400
        ) from error
    try:
        yield form
    finally:
        await form.close()


def attachment_name(filename: str | None) -> str:
    # Имя используется только как подпись, файл на диск под этим именем не пишется.
    name = PurePosixPath((filename or "").replace("\\", "/")).name
    return " ".join(name.split()) or "без имени"


def validate_upload(upload: UploadFile) -> str:
    name = attachment_name(upload.filename)
    extension = PurePosixPath(name).suffix.lower()
    mime = (upload.content_type or "").split(";", 1)[0].strip().lower()
    if extension not in ALLOWED_TYPES or mime != ALLOWED_TYPES[extension]:
        raise AttachmentError(
            f"Неподдерживаемый тип файла «{name}» или несоответствие content-type. "
            "Разрешены JPG, JPEG, PNG, PDF, XLSX, XLS и DOCX.",
            415,
        )
    if upload.size is not None and upload.size > MAX_FILE_SIZE:
        raise AttachmentError(f"Файл «{name}» превышает лимит 10 МБ.", 413)
    return extension


def _row_text(values) -> str:
    cells = ["" if value is None else str(value).strip() for value in values]
    # Внутренние пустые колонки сохраняют соответствие заголовкам.
    while cells and not cells[-1]:
        cells.pop()
    return " | ".join(cells) if any(cells) else ""


def _document_lines(data: bytes, extension: str) -> Iterator[str]:
    stream = BytesIO(data)
    if extension == ".pdf":
        reader = PdfReader(stream)
        for page in reader.pages[:10]:
            yield page.extract_text() or ""
        if len(reader.pages) > 10:
            yield "[Документ обрезан: прочитаны только первые 10 страниц PDF.]"
    elif extension == ".xlsx":
        workbook = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    yield _row_text(row)
        finally:
            workbook.close()
    elif extension == ".xls":
        with xlrd.open_workbook(file_contents=data, on_demand=True) as workbook:
            for sheet in workbook.sheets():
                for row_index in range(sheet.nrows):
                    values = sheet.row_values(row_index)
                    yield _row_text(
                        int(value) if isinstance(value, float) and value.is_integer() else value
                        for value in values
                    )
    elif extension == ".docx":
        document = Document(stream)
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                for row in block.rows:
                    yield _row_text(cell.text for cell in row.cells)
            else:
                yield block.text


def extract_text(data: bytes, extension: str) -> str:
    """Извлекает не более 8000 символов, включая пометку об обрезке."""
    text = ""
    with closing(_document_lines(data, extension)) as lines:
        for line in lines:
            line = line.strip()
            if not line:
                continue
            text += ("\n" if text else "") + line
            if len(text) > MAX_TEXT_LENGTH:
                return text[:MAX_TEXT_LENGTH - len(TRUNCATION_NOTICE)] + TRUNCATION_NOTICE
    return text


def prepare_attachments(uploads: list[UploadFile]) -> list[dict]:
    """Вызывается в пуле потоков; файлы закрывает владелец multipart-формы."""
    if len(uploads) > MAX_FILES:
        raise AttachmentError("Можно прикрепить не более 3 файлов.", 413)
    extensions = [validate_upload(upload) for upload in uploads]
    parts = []
    for upload, extension in zip(uploads, extensions):
        name = attachment_name(upload.filename)
        data = upload.file.read(MAX_FILE_SIZE + 1)
        if len(data) > MAX_FILE_SIZE:
            raise AttachmentError(f"Файл «{name}» превышает лимит 10 МБ.", 413)
        if not data:
            raise AttachmentError(f"Файл «{name}» пуст. Прикрепите файл с содержимым.")
        if extension in (".jpg", ".jpeg", ".png"):
            signature = b"\x89PNG\r\n\x1a\n" if extension == ".png" else b"\xff\xd8\xff"
            if not data.startswith(signature):
                raise AttachmentError(f"Не удалось прочитать изображение «{name}»: файл повреждён.")
            encoded = base64.b64encode(data).decode("ascii")
            parts.extend([
                {"type": "text", "text": f"Фото товара «{name}»:"},
                {"type": "image_url", "image_url": {"url": f"data:{ALLOWED_TYPES[extension]};base64,{encoded}"}},
            ])
        else:
            try:
                text = extract_text(data, extension)
            except Exception as error:
                raise AttachmentError(
                    f"Не удалось прочитать файл «{name}». Проверьте, что он не повреждён "
                    "и не защищён паролем."
                ) from error
            if not text:
                text = (
                    "[Текст не найден. Если это скан, попроси прислать страницы как фото "
                    "JPG/PNG или документ с текстовым слоем.]"
                )
            parts.append({"type": "text", "text": f"Содержимое вложения «{name}»:\n{text}"})
    return parts
